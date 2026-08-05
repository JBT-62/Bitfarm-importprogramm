#!/usr/bin/env python3
"""
export_angebote_pruefung.py
Kunden-Angebote (ANGAUFGUT WHERE ANGEBOT=1) mit Auftragswert
< 50.000 € ODER > 150.000 € – Artikel lesen, Excel erstellen.

Ausführen:
    python 99_scripts/export_angebote_pruefung.py

Ergebnis:
    Angebote_Pruefung_JJJJMMTT.xlsx
"""

import os, sys
from datetime import date

try:
    from dotenv import load_dotenv
    load_dotenv()
except ImportError:
    pass

try:
    import pyodbc
except ImportError:
    sys.exit("FEHLER: pyodbc fehlt.  pip install pyodbc")

try:
    import openpyxl
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    from openpyxl.utils import get_column_letter
except ImportError:
    sys.exit("FEHLER: openpyxl fehlt.  pip install openpyxl")


LIMIT_KLEIN = 50_000
LIMIT_GROSS = 150_000


# ── Verbindung ────────────────────────────────────────────────────────────────

def ee_connect():
    driver   = os.getenv("EV_DRIVER",   "ODBC Driver 18 for SQL Server")
    server   = os.getenv("EV_SERVER",   "192.168.120.234")
    database = os.getenv("EV_DATABASE", "KuF")
    user     = os.getenv("EV_USER",     "excel_kuf_readonly")
    password = os.getenv("EV_PASSWORD", "")
    conn_str = (
        f"DRIVER={{{driver}}};"
        f"SERVER={server},1433;"
        f"DATABASE={database};"
        f"UID={user};"
        f"PWD={password};"
        f"TrustServerCertificate=yes;"
    )
    conn = pyodbc.connect(conn_str, timeout=10)
    print(f"[OK] Verbunden mit {server}/{database}")
    return conn


def q(conn, sql, params=None):
    cur = conn.cursor()
    cur.execute(sql, params or [])
    cols = [c[0] for c in cur.description]
    return [dict(zip(cols, r)) for r in cur.fetchall()]


def spalten_set(conn, tabelle):
    rows = q(conn, """
        SELECT COLUMN_NAME, DATA_TYPE
        FROM INFORMATION_SCHEMA.COLUMNS
        WHERE TABLE_NAME = ?
        ORDER BY ORDINAL_POSITION
    """, [tabelle])
    return ({r['COLUMN_NAME'].upper(): r['DATA_TYPE'] for r in rows},
            [r['COLUMN_NAME'] for r in rows])


def erste(kandidaten, vorhanden_set):
    for k in kandidaten:
        if k.upper() in vorhanden_set:
            return k
    return None


# ── Schema erkunden ───────────────────────────────────────────────────────────

def erkunde(conn):
    print("\n[1] Erkunde ANGAUFGUT / ANGAUFPOS …")

    h_t, h_c = spalten_set(conn, "ANGAUFGUT")
    p_t, p_c = spalten_set(conn, "ANGAUFPOS")

    print(f"    ANGAUFGUT  : {', '.join(h_c)}")
    print(f"    ANGAUFPOS  : {', '.join(p_c)}")

    h_set = set(h_t.keys())
    p_set = set(p_t.keys())

    # Kopf-Wert
    wert_kopf = erste(
        ["GESAMTBETRAG", "NETTOBETRAG", "AUFTRAGSWERT", "NETTO",
         "BRUTTOBETRAG", "BETRAG", "SUMME", "WERT"],
        h_set
    )
    # Positionen
    art_col  = erste(["ARTNR1", "ARTNR", "ARTIKELNR"],           p_set)
    bez_col  = erste(["ABEZ1", "ARTBEZ", "BEZEICHNUNG", "BEZ"],  p_set)
    mng_col  = erste(["MENGE", "ANZAHL"],                         p_set)
    vk_col   = erste(["VKPR1", "VK1", "VKPREIS", "PREIS", "EP"], p_set)
    ek_col   = erste(["EKPR", "EK1", "EKPREIS"],                  p_set)
    pos_wert = erste(["GESAMTPREIS", "POSITIONSWERT", "BETRAG",
                      "ZEILENPREIS", "POSBET", "NETTO"],           p_set)
    # FK Position → Kopf
    fk = erste(
        ["ANGAUFGUTLFDNR", "ANGAUFLFDNR", "LFDANGAUFGUT",
         "KOPFLFDNR",      "HEADLFDNR",   "BELEGLFDNR",
         "ANGAUFTLFDNR"],
        p_set
    )
    # Kopf-Infos
    knd_col = erste(["KNDNR", "ADRNR", "KUNDENNR"], h_set)
    dat_col = erste(["ERFASSDATUM", "DATUM", "BELEGDATUM", "ERFDATUM"], h_set)
    sta_col = erste(["STATUS", "ERLEDIGT"], h_set)
    vtr_col = erste(["VERMITNR", "VERTRETERNR", "SACHBEARBEITNR", "BETREUNR"], h_set)

    info = dict(
        wert_kopf=wert_kopf, art_col=art_col, bez_col=bez_col,
        mng_col=mng_col,     vk_col=vk_col,   ek_col=ek_col,
        pos_wert=pos_wert,   fk=fk,
        knd_col=knd_col,     dat_col=dat_col,
        sta_col=sta_col,     vtr_col=vtr_col,
    )
    print("    Erkannte Felder:")
    for k, v in info.items():
        print(f"      {k:<12} = {v}")
    return info


# ── Angebote laden ────────────────────────────────────────────────────────────

def lade_angebote(conn, info):
    print(f"\n[2] Lade Kunden-Angebote (ANGEBOT=1), Wert < {LIMIT_KLEIN:,.0f} € oder > {LIMIT_GROSS:,.0f} € …")

    wk  = info["wert_kopf"]
    fk  = info["fk"]
    art = info["art_col"]
    bez = info["bez_col"]
    mng = info["mng_col"]
    vk  = info["vk_col"]
    ek  = info["ek_col"]
    pw  = info["pos_wert"]
    knd = info["knd_col"]
    dat = info["dat_col"]
    sta = info["sta_col"]
    vtr = info["vtr_col"]

    if not fk:
        print("    FEHLER: FK-Spalte ANGAUFPOS→ANGAUFGUT nicht erkannt.")
        print(f"    Bekannte ANGAUFPOS-Spalten oben prüfen.")
        return [], []

    # Wert-Berechnung
    if wk:
        where_w = f"h.{wk} < {LIMIT_KLEIN} OR h.{wk} > {LIMIT_GROSS}"
        w_sel   = f"h.{wk} AS Auftragswert"
    elif vk and mng:
        sub = f"(SELECT COALESCE(SUM(pp.{mng}*pp.{vk}),0) FROM ANGAUFPOS pp WHERE pp.{fk}=h.LFDNR)"
        where_w = f"{sub} < {LIMIT_KLEIN} OR {sub} > {LIMIT_GROSS}"
        w_sel   = f"{sub} AS Auftragswert"
    elif pw:
        sub = f"(SELECT COALESCE(SUM(pp.{pw}),0) FROM ANGAUFPOS pp WHERE pp.{fk}=h.LFDNR)"
        where_w = f"{sub} < {LIMIT_KLEIN} OR {sub} > {LIMIT_GROSS}"
        w_sel   = f"{sub} AS Auftragswert"
    else:
        print("    WARNUNG: Keine Preisspalte – lade alle Angebote.")
        where_w = "1=1"
        w_sel   = "NULL AS Auftragswert"

    felder = ["h.LFDNR AS Angebot_ID", "h.ANGEBOT", "h.AUFTRAG"]
    if knd: felder.append(f"h.{knd} AS Kunden_Nr")
    if dat: felder.append(f"h.{dat} AS Datum")
    if sta: felder.append(f"h.{sta} AS Status")
    if vtr: felder.append(f"h.{vtr} AS Vertreter_Nr")
    felder.append(w_sel)

    sql = f"""
        SELECT {', '.join(felder)}
        FROM ANGAUFGUT h
        WHERE h.ANGEBOT = 1
          AND ({where_w})
        ORDER BY h.LFDNR DESC
    """

    try:
        angebote = q(conn, sql)
    except Exception as e:
        print(f"    Hauptabfrage fehlgeschlagen: {e}")
        print("    Fallback: alle Angebote laden …")
        sql_fb = "SELECT TOP 1000 * FROM ANGAUFGUT WHERE ANGEBOT=1 ORDER BY LFDNR DESC"
        angebote = q(conn, sql_fb)
        for a in angebote:
            a.setdefault("Angebot_ID", a.get("LFDNR"))
            a.setdefault("Auftragswert", 0)

    print(f"    {len(angebote)} Angebote gefunden")
    if not angebote:
        return [], []

    # ── Positionen ────────────────────────────────────────────────────────────
    ids = [str(int(a["Angebot_ID"])) for a in angebote if a.get("Angebot_ID") is not None]
    if not ids:
        return angebote, []

    id_liste = ", ".join(ids[:500])

    pos_felder = [f"p.{fk} AS Angebot_ID"]
    if art: pos_felder.append(f"p.{art} AS Artikelnummer")
    if bez: pos_felder.append(f"p.{bez} AS Bezeichnung")
    if mng: pos_felder.append(f"p.{mng} AS Menge")
    if vk:  pos_felder.append(f"p.{vk} AS VK_Preis")
    if ek:  pos_felder.append(f"p.{ek} AS EK_Preis")
    if pw:  pos_felder.append(f"p.{pw} AS Positionswert")

    # Langtext
    if art:
        sel_lang = ", ".join(pos_felder) + """,
            al.CONTENT AS Langtext,
            ar.EKPR    AS EK_Stamm,
            ar.VKPR1   AS VK_Stamm,
            ar.MENGENSCHL AS Einheit"""
        join_lang = f"""
            LEFT JOIN ARTIKEL ar ON ar.ARTNR1 = p.{art}
            LEFT JOIN (
                SELECT LFDARTNR, MAX(CONTENT) AS CONTENT
                FROM ARTIKELLONG GROUP BY LFDARTNR
            ) al ON al.LFDARTNR = ar.LFDNR"""
    else:
        sel_lang = ", ".join(pos_felder)
        join_lang = ""

    order_p = f"p.{fk}" + (f", p.{art}" if art else "")

    sql_pos = f"""
        SELECT {sel_lang}
        FROM ANGAUFPOS p
        {join_lang}
        WHERE p.{fk} IN ({id_liste})
        ORDER BY {order_p}
    """

    try:
        positionen = q(conn, sql_pos)
    except Exception as e:
        print(f"    JOIN fehlgeschlagen ({e}), einfacher Fallback …")
        sql_pos2 = f"""
            SELECT {', '.join(pos_felder)}
            FROM ANGAUFPOS p
            WHERE p.{fk} IN ({id_liste})
            ORDER BY {order_p}
        """
        try:
            positionen = q(conn, sql_pos2)
        except Exception as e2:
            print(f"    Auch einfacher Fallback fehlgeschlagen: {e2}")
            positionen = []

    print(f"    {len(positionen)} Positionen geladen")

    pos_map = {}
    for p in positionen:
        pos_map.setdefault(p.get("Angebot_ID"), []).append(p)
    for a in angebote:
        a["_positionen"] = pos_map.get(a.get("Angebot_ID"), [])

    return angebote, positionen


# ── Excel-Stile ───────────────────────────────────────────────────────────────

BLAU_D = "1B4F8A"
BLAU_H = "D6E4F7"
GELB   = "FFF2CC"
GRUEN  = "E2EFDA"
ROT_H  = "FCE4D6"
WEISS  = "FFFFFF"

def stil(cell, fett=False, bg=WEISS, farbe="000000", groesse=9, mitte=False, zahl=False):
    cell.font      = Font(name="Arial", bold=fett, color=farbe, size=groesse)
    cell.fill      = PatternFill("solid", fgColor=bg)
    ha = "center" if mitte else ("right" if zahl else "left")
    cell.alignment = Alignment(horizontal=ha, vertical="center", wrap_text=False)
    cell.border    = Border(
        bottom=Side(style="thin", color="DDDDDD"),
        right=Side(style="thin", color="DDDDDD"),
    )

def kopfzeile(ws, zeile, texte, bg=BLAU_D):
    for col, txt in enumerate(texte, 1):
        c = ws.cell(row=zeile, column=col, value=txt)
        stil(c, fett=True, bg=bg, farbe="FFFFFF", groesse=9, mitte=True)
    ws.row_dimensions[zeile].height = 22

def datenzeile(ws, zeile, werte, bg=WEISS):
    for col, wert in enumerate(werte, 1):
        c = ws.cell(row=zeile, column=col, value=wert)
        is_z = isinstance(wert, (int, float))
        stil(c, bg=bg, zahl=is_z)
        if is_z and col >= 4:
            c.number_format = '#,##0.00 €'

def breiten(ws, bmap):
    for col, b in bmap.items():
        ws.column_dimensions[col].width = b


# ── Blatt 1: Angebots-Übersicht ───────────────────────────────────────────────

def blatt_angebote(wb, angebote):
    ws = wb.create_sheet("Angebote")
    ws.sheet_view.showGridLines = False
    ws["B1"].value = (
        f"Kunden-Angebote – Wert < {LIMIT_KLEIN:,.0f} € oder > {LIMIT_GROSS:,.0f} €"
    )
    ws["B1"].font = Font(name="Arial", bold=True, size=13, color=BLAU_D)
    ws.row_dimensions[1].height = 30

    if not angebote:
        ws["B3"].value = "(keine Angebote gefunden)"; return

    keys = [k for k in angebote[0] if k != "_positionen"]
    kopfzeile(ws, 3, keys)

    for i, a in enumerate(angebote, 4):
        wert = a.get("Auftragswert") or 0
        try:    wert_f = float(wert)
        except: wert_f = 0
        bg = ROT_H if wert_f > LIMIT_GROSS else (BLAU_H if wert_f < LIMIT_KLEIN else WEISS)
        datenzeile(ws, i, [a.get(k) for k in keys], bg=bg)

    leg = len(angebote) + 5
    ws.cell(row=leg,   column=2, value="Rot  = Großauftrag > 150 k €").font = Font(name="Arial", size=8, color="888888")
    ws.cell(row=leg+1, column=2, value="Blau = Kleinauftrag < 50 k €").font = Font(name="Arial", size=8, color="888888")

    ws.auto_filter.ref = f"B3:{get_column_letter(len(keys)+1)}{len(angebote)+3}"
    ws.freeze_panes    = ws.cell(row=4, column=1)
    breiten(ws, {"A": 3, "B": 14, "C": 10, "D": 10, "E": 16, "F": 18, "G": 10, "H": 16, "I": 18})


# ── Blatt 2: Alle Positionen ──────────────────────────────────────────────────

def blatt_positionen(wb, positionen):
    ws = wb.create_sheet("Artikel-Positionen")
    ws.sheet_view.showGridLines = False
    ws["B1"].value = "Angebotspositionen – alle Artikel"
    ws["B1"].font  = Font(name="Arial", bold=True, size=13, color=BLAU_D)
    ws.row_dimensions[1].height = 30

    if not positionen:
        ws["B3"].value = "(keine Positionen)"; return

    keys = list(positionen[0].keys())
    kopfzeile(ws, 3, keys)
    for i, p in enumerate(positionen, 4):
        bg = BLAU_H if i % 2 == 0 else WEISS
        datenzeile(ws, i, [p.get(k) for k in keys], bg=bg)

    ws.auto_filter.ref = f"A3:{get_column_letter(len(keys))}{len(positionen)+3}"
    ws.freeze_panes    = ws.cell(row=4, column=1)
    breiten(ws, {"A": 14, "B": 18, "C": 32, "D": 10, "E": 14, "F": 14, "G": 16, "H": 12, "I": 14, "J": 40})


# ── Blatt 3: Artikel-Zusammenfassung ─────────────────────────────────────────

def blatt_artikel(wb, positionen):
    ws = wb.create_sheet("Artikel-Zusammenfassung")
    ws.sheet_view.showGridLines = False
    ws["B1"].value = "Artikel-Zusammenfassung (dedup, sortiert nach Häufigkeit)"
    ws["B1"].font  = Font(name="Arial", bold=True, size=13, color=BLAU_D)
    ws.row_dimensions[1].height = 30

    art_map = {}
    for p in positionen:
        key = str(p.get("Artikelnummer") or p.get("Bezeichnung") or "–")
        if key not in art_map:
            art_map[key] = {
                "Artikelnummer":   p.get("Artikelnummer"),
                "Bezeichnung":     p.get("Bezeichnung"),
                "Einheit":         p.get("Einheit"),
                "EK_Stamm €":     p.get("EK_Stamm"),
                "VK_Stamm €":     p.get("VK_Stamm"),
                "Anzahl_Angebote": 0,
                "Gesamt_Menge":   0,
                "Langtext":        p.get("Langtext"),
            }
        art_map[key]["Anzahl_Angebote"] += 1
        try:    art_map[key]["Gesamt_Menge"] += float(p.get("Menge") or 0)
        except: pass

    sortiert = sorted(art_map.values(), key=lambda x: x["Anzahl_Angebote"], reverse=True)

    headers = ["Artikelnummer", "Bezeichnung", "Einheit",
               "EK-Preis €", "VK-Preis €",
               "Anzahl Angebote", "Gesamt-Menge",
               "✓ Korrekt?", "Anmerkung", "Langtext"]
    kopfzeile(ws, 3, headers)

    for i, art in enumerate(sortiert, 4):
        bg = GRUEN if i % 2 == 0 else WEISS
        werte = [
            art.get("Artikelnummer"), art.get("Bezeichnung"),
            art.get("Einheit"),
            art.get("EK_Stamm €"),   art.get("VK_Stamm €"),
            art.get("Anzahl_Angebote"), art.get("Gesamt_Menge"),
            "", "",                  # Freigabe / Anmerkung
            art.get("Langtext"),
        ]
        datenzeile(ws, i, werte, bg=bg)
        for col in [8, 9]:
            c = ws.cell(row=i, column=col)
            c.fill = PatternFill("solid", fgColor=GELB)

    ws.auto_filter.ref = f"A3:{get_column_letter(len(headers))}{len(sortiert)+3}"
    ws.freeze_panes    = ws.cell(row=4, column=1)
    breiten(ws, {"A": 18, "B": 32, "C": 10, "D": 14, "E": 14,
                 "F": 16, "G": 14, "H": 14, "I": 28, "J": 42})


# ── Blatt 4: Freigabe ─────────────────────────────────────────────────────────

def blatt_freigabe(wb, heute, n_angebote, n_artikel):
    ws = wb.create_sheet("Freigabe")
    ws.sheet_view.showGridLines = False
    ws["B2"].value = "Freigabe – Kunden-Angebote Artikel-Prüfung"
    ws["B2"].font  = Font(name="Arial", bold=True, size=14, color=BLAU_D)
    ws.row_dimensions[2].height = 32
    ws["B3"].value = (
        f"Exportiert {heute.strftime('%d.%m.%Y')}  ·  "
        f"{n_angebote} Angebote  ·  {n_artikel} eindeutige Artikel"
    )
    ws["B3"].font = Font(name="Arial", size=9, color="888888")

    for zeile, label, vor in [
        (6,  "Prüfer (Name)",                 ""),
        (7,  "Abteilung",                     "Vertrieb"),
        (8,  "Datum",                         heute.strftime("%d.%m.%Y")),
        (9,  "Angebots-Liste geprüft?",       "☐  Ja    ☐  Nein"),
        (10, "Artikel vollständig & korrekt?","☐  Ja    ☐  Nein"),
        (12, "Gesamtfreigabe",                "☐  Freigegeben    ☐  Korrekturen nötig"),
        (14, "Unterschrift",                  ""),
    ]:
        ws.cell(row=zeile, column=2, value=label).font = Font(name="Arial", bold=True, size=10)
        c = ws.cell(row=zeile, column=3, value=vor)
        c.font = Font(name="Arial", size=10)
        c.fill = PatternFill("solid", fgColor=GELB)
        ws.row_dimensions[zeile].height = 24
        if "Unterschrift" in label:
            ws.row_dimensions[zeile].height = 42
            c.value  = ""
            c.border = Border(bottom=Side(style="medium", color="333333"))

    ws["B16"].value = "Anmerkungen:"
    ws["B16"].font  = Font(name="Arial", bold=True, size=10)
    ws.merge_cells("B17:F22")
    anm = ws["B17"]
    anm.fill      = PatternFill("solid", fgColor=GELB)
    anm.alignment = Alignment(vertical="top", wrap_text=True)
    anm.border    = Border(
        top=Side(style="thin"), bottom=Side(style="thin"),
        left=Side(style="thin"), right=Side(style="thin"),
    )
    for r in range(17, 23):
        ws.row_dimensions[r].height = 18
    breiten(ws, {"A": 3, "B": 38, "C": 50})


# ── Hauptprogramm ─────────────────────────────────────────────────────────────

def main():
    heute = date.today()
    conn  = ee_connect()

    info = erkunde(conn)
    angebote, positionen = lade_angebote(conn, info)
    conn.close()

    if not angebote:
        print("\n[!] Keine Angebote gefunden.")
        print("    Mögliche Ursachen:")
        print("    - ANGEBOT-Flag nicht = 1 (Spaltenname prüfen)")
        print("    - Alle Angebote liegen zwischen 50 k und 150 k €")
        print("    - FK-Spalte ANGAUFPOS→ANGAUFGUT nicht erkannt")
        return

    n_art = len({p.get("Artikelnummer") or p.get("Bezeichnung") for p in positionen})

    print(f"\n[3] Erstelle Excel …")
    wb = openpyxl.Workbook()
    wb.remove(wb.active)

    blatt_angebote(wb, angebote)
    blatt_positionen(wb, positionen)
    blatt_artikel(wb, positionen)
    blatt_freigabe(wb, heute, len(angebote), n_art)

    datei = f"Angebote_Pruefung_{heute.strftime('%Y%m%d')}.xlsx"
    wb.save(datei)
    print(f"\n[OK] Gespeichert: {datei}")
    print(f"     {len(angebote)} Angebote  ·  {len(positionen)} Positionen  ·  {n_art} eindeutige Artikel")
    print(f"     Bitte an Vertrieb / Produktionsleiter weiterleiten.")


if __name__ == "__main__":
    main()
