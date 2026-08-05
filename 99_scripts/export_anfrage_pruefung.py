#!/usr/bin/env python3
"""
export_anfrage_pruefung.py
Findet CRM-Anfragen mit Auftragswert < 50.000 € ODER > 150.000 €,
liest die angeforderten Artikel und erstellt eine Excel-Prüfdatei.

Ausführen:
    python 99_scripts/export_anfrage_pruefung.py

Ergebnis:
    Anfragen_Pruefung_JJJJMMTT.xlsx
"""

import os, sys
from datetime import date

try:
    from dotenv import load_dotenv
    load_dotenv()
except ImportError:
    pass

try:
    import pyodbc
except ImportError:
    sys.exit("FEHLER: pyodbc fehlt.  pip install pyodbc")

try:
    import openpyxl
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    from openpyxl.utils import get_column_letter
except ImportError:
    sys.exit("FEHLER: openpyxl fehlt.  pip install openpyxl")


LIMIT_KLEIN = 50_000     # < diese Grenze
LIMIT_GROSS = 150_000    # > diese Grenze


# ── Verbindung ────────────────────────────────────────────────────────────────

def ee_connect():
    driver   = os.getenv("EV_DRIVER",   "ODBC Driver 18 for SQL Server")
    server   = os.getenv("EV_SERVER",   "192.168.120.234")
    database = os.getenv("EV_DATABASE", "KuF")
    user     = os.getenv("EV_USER",     "excel_kuf_readonly")
    password = os.getenv("EV_PASSWORD", "")
    conn_str = (
        f"DRIVER={{{driver}}};"
        f"SERVER={server},1433;"
        f"DATABASE={database};"
        f"UID={user};"
        f"PWD={password};"
        f"TrustServerCertificate=yes;"
    )
    conn = pyodbc.connect(conn_str, timeout=10)
    print(f"[OK] Verbunden mit {server}/{database}")
    return conn


def q(conn, sql, params=None):
    cur = conn.cursor()
    cur.execute(sql, params or [])
    cols = [c[0] for c in cur.description]
    return [dict(zip(cols, r)) for r in cur.fetchall()]


def spalten(conn, tabelle):
    rows = q(conn, """
        SELECT COLUMN_NAME, DATA_TYPE
        FROM INFORMATION_SCHEMA.COLUMNS
        WHERE TABLE_NAME = ?
        ORDER BY ORDINAL_POSITION
    """, [tabelle])
    return {r['COLUMN_NAME'].upper(): r['DATA_TYPE'] for r in rows}, \
           [r['COLUMN_NAME'] for r in rows]


def erste(kandidaten, vorhanden):
    """Ersten Kandidaten zurückgeben der in vorhanden (set von UPPER) ist."""
    for k in kandidaten:
        if k.upper() in vorhanden:
            return k
    return None


# ── Schema-Erkundung ─────────────────────────────────────────────────────────

def erkunde_schema(conn):
    print("\n[1] Erkunde ANFRAGE-Schema …")
    anf_t, anf_c = spalten(conn, "ANFRAGE")
    pos_t, pos_c = spalten(conn, "ANFRAGEPOS")

    print(f"    ANFRAGE    : {', '.join(anf_c)}")
    print(f"    ANFRAGEPOS : {', '.join(pos_c)}")

    anf_set = set(anf_t.keys())
    pos_set = set(pos_t.keys())

    # ── Kopftabelle: Wertspalte ───────────────────────────────────────────────
    wert_kopf = erste(
        ["GESAMTBETRAG", "NETTOBETRAG", "AUFTRAGSWERT", "BETRAG",
         "SUMME", "BRUTTOBETRAG", "WERT", "GESAMTWERT"],
        anf_set
    )

    # ── Positionstabelle: Preisspalten ────────────────────────────────────────
    art_col  = erste(["ARTNR1", "ARTNR", "ARTIKELNR"],        pos_set)
    bez_col  = erste(["ABEZ1", "ARTBEZ", "BEZEICHNUNG", "BEZ"], pos_set)
    mng_col  = erste(["MENGE", "ANZAHL"],                      pos_set)
    vk_col   = erste(["VKPR1", "VK1", "VKPREIS", "PREIS", "EP"], pos_set)
    ek_col   = erste(["EKPR", "EK1", "EKPREIS", "EP_EK"],     pos_set)
    pos_wert = erste(["GESAMTPREIS", "POSITIONSWERT", "BETRAG",
                      "ZEILENPREIS", "POSBET"],                 pos_set)

    # ── FK Kopf → Position ────────────────────────────────────────────────────
    fk = erste(
        ["ANFRAGELFDNR", "ANFRAGEID", "KOPFLFDNR",
         "LFDANFRAGE",   "HEADLFDNR", "BELEGLFDNR"],
        pos_set
    )

    # ── Kopf-Infospalten ──────────────────────────────────────────────────────
    knd_col  = erste(["KNDNR", "ADRNR", "KUNDENNR"], anf_set)
    dat_col  = erste(["DATUM", "ERFASSDATUM", "BELEGDATUM", "ERFDATUM"], anf_set)
    sta_col  = erste(["STATUS", "ERLEDIGT"], anf_set)
    nr_col   = erste(["BELEGNR", "ANFRAGNR", "LFDNR"], anf_set)
    name_col = erste(["NAME1", "MATCHCODE", "SUCHBEGRIFF"], anf_set)

    info = {
        "wert_kopf": wert_kopf,
        "art_col":   art_col,
        "bez_col":   bez_col,
        "mng_col":   mng_col,
        "vk_col":    vk_col,
        "ek_col":    ek_col,
        "pos_wert":  pos_wert,
        "fk":        fk,
        "knd_col":   knd_col,
        "dat_col":   dat_col,
        "sta_col":   sta_col,
        "nr_col":    nr_col,
        "name_col":  name_col,
    }
    print("    Erkannte Felder:")
    for k, v in info.items():
        print(f"      {k:<12} = {v}")

    return info


# ── Anfragen laden ────────────────────────────────────────────────────────────

def lade_anfragen(conn, info):
    print(f"\n[2] Suche Anfragen  < {LIMIT_KLEIN:,.0f} € ODER > {LIMIT_GROSS:,.0f} € …")

    fk       = info["fk"]
    art      = info["art_col"]
    bez      = info["bez_col"]
    mng      = info["mng_col"]
    vk       = info["vk_col"]
    ek       = info["ek_col"]
    pos_wert = info["pos_wert"]
    knd      = info["knd_col"]
    dat      = info["dat_col"]
    sta      = info["sta_col"]
    nr       = info["nr_col"]
    name     = info["name_col"]
    wk       = info["wert_kopf"]

    if not fk:
        print("    FEHLER: FK-Spalte ANFRAGE→ANFRAGEPOS nicht erkannt.")
        print("    Bitte ANFRAGEPOS-Spalten manuell prüfen (suche_pia_anfrage2.py).")
        return [], []

    # Auftragswert: entweder aus Kopftabelle oder per Summe aus Positionen
    if wk:
        # Wert direkt im Kopf
        where_wert = f"h.{wk} < {LIMIT_KLEIN} OR h.{wk} > {LIMIT_GROSS}"
        wert_select = f"h.{wk} AS Auftragswert"
    elif vk and mng:
        # Wert als Subquery aus ANFRAGEPOS berechnen
        where_wert = f"""
            (SELECT COALESCE(SUM(pp.{mng} * pp.{vk}), 0)
             FROM ANFRAGEPOS pp WHERE pp.{fk} = h.LFDNR)
            < {LIMIT_KLEIN}
            OR
            (SELECT COALESCE(SUM(pp.{mng} * pp.{vk}), 0)
             FROM ANFRAGEPOS pp WHERE pp.{fk} = h.LFDNR)
            > {LIMIT_GROSS}
        """
        wert_select = f"""
            (SELECT COALESCE(SUM(pp.{mng} * pp.{vk}), 0)
             FROM ANFRAGEPOS pp WHERE pp.{fk} = h.LFDNR) AS Auftragswert
        """
    elif pos_wert:
        where_wert = f"""
            (SELECT COALESCE(SUM(pp.{pos_wert}), 0)
             FROM ANFRAGEPOS pp WHERE pp.{fk} = h.LFDNR)
            < {LIMIT_KLEIN}
            OR
            (SELECT COALESCE(SUM(pp.{pos_wert}), 0)
             FROM ANFRAGEPOS pp WHERE pp.{fk} = h.LFDNR)
            > {LIMIT_GROSS}
        """
        wert_select = f"""
            (SELECT COALESCE(SUM(pp.{pos_wert}), 0)
             FROM ANFRAGEPOS pp WHERE pp.{fk} = h.LFDNR) AS Auftragswert
        """
    else:
        print("    WARNUNG: Keine Wert-/Preisspalte erkannt – lade alle Anfragen.")
        where_wert = "1=1"
        wert_select = "NULL AS Auftragswert"

    # Kopf-SELECT aufbauen
    kopf_felder = ["h.LFDNR AS Anfrage_ID"]
    if nr and nr != "LFDNR":    kopf_felder.append(f"h.{nr} AS Belegnummer")
    if knd:                      kopf_felder.append(f"h.{knd} AS Kunden_Nr")
    if name:                     kopf_felder.append(f"h.{name} AS Kunde")
    if dat:                      kopf_felder.append(f"h.{dat} AS Datum")
    if sta:                      kopf_felder.append(f"h.{sta} AS Status")
    kopf_felder.append(wert_select)

    sql_anfragen = f"""
        SELECT {', '.join(kopf_felder)}
        FROM ANFRAGE h
        WHERE ({where_wert})
        ORDER BY h.LFDNR DESC
    """

    try:
        anfragen = q(conn, sql_anfragen)
    except Exception as e:
        print(f"    Hauptabfrage fehlgeschlagen: {e}")
        print("    Versuche einfachere Abfrage (alle Anfragen, ohne Wertfilter) …")
        sql_einfach = "SELECT TOP 200 * FROM ANFRAGE ORDER BY LFDNR DESC"
        anfragen = q(conn, sql_einfach)
        # Nachfiltern
        result = []
        for a in anfragen:
            wert = a.get("Auftragswert") or a.get(wk or "") or 0
            try:
                wert = float(wert or 0)
            except (TypeError, ValueError):
                wert = 0
            if wert < LIMIT_KLEIN or wert > LIMIT_GROSS:
                result.append(a)
        anfragen = result

    print(f"    {len(anfragen)} Anfragen gefunden")

    if not anfragen:
        return [], []

    # ── Positionen laden ──────────────────────────────────────────────────────
    ids = [str(int(a["Anfrage_ID"])) for a in anfragen if a.get("Anfrage_ID")]
    id_liste = ", ".join(ids[:200])

    pos_felder = [f"p.{fk} AS Anfrage_ID"]
    if art:      pos_felder.append(f"p.{art} AS Artikelnummer")
    if bez:      pos_felder.append(f"p.{bez} AS Bezeichnung")
    if mng:      pos_felder.append(f"p.{mng} AS Menge")
    if vk:       pos_felder.append(f"p.{vk} AS VK_Preis")
    if ek:       pos_felder.append(f"p.{ek} AS EK_Preis")
    if pos_wert: pos_felder.append(f"p.{pos_wert} AS Positionswert")
    # Langtext aus ARTIKELLONG
    if art:
        pos_felder_sql = ", ".join(pos_felder) + f"""
            , al.CONTENT AS Langtext
        """
        join_lang = f"""
            LEFT JOIN (
                SELECT LFDARTNR, MAX(CONTENT) AS CONTENT
                FROM ARTIKELLONG GROUP BY LFDARTNR
            ) al ON al.LFDARTNR = (
                SELECT TOP 1 LFDNR FROM ARTIKEL WHERE ARTNR1 = p.{art}
            )
        """
    else:
        pos_felder_sql = ", ".join(pos_felder)
        join_lang = ""

    sql_pos = f"""
        SELECT {pos_felder_sql}
        FROM ANFRAGEPOS p
        {join_lang}
        WHERE p.{fk} IN ({id_liste})
        ORDER BY p.{fk}, p.LFDNR
    """

    try:
        positionen = q(conn, sql_pos)
    except Exception as e:
        print(f"    Langtext-JOIN fehlgeschlagen ({e}), versuche ohne …")
        sql_pos_einfach = f"""
            SELECT {', '.join(pos_felder)}
            FROM ANFRAGEPOS p
            WHERE p.{fk} IN ({id_liste})
            ORDER BY p.{fk}, p.LFDNR
        """
        positionen = q(conn, sql_pos_einfach)

    print(f"    {len(positionen)} Positionen geladen")

    # Positionen an Anfrage anhängen
    pos_map = {}
    for p in positionen:
        aid = p.get("Anfrage_ID")
        pos_map.setdefault(aid, []).append(p)

    for a in anfragen:
        a["_positionen"] = pos_map.get(a.get("Anfrage_ID"), [])

    return anfragen, positionen


# ── Excel-Stile ───────────────────────────────────────────────────────────────

BLAU_D = "1B4F8A"
BLAU_H = "D6E4F7"
GELB   = "FFF2CC"
GRUEN  = "E2EFDA"
ROT_H  = "FCE4D6"
WEISS  = "FFFFFF"

def stil(cell, fett=False, bg=WEISS, farbe="000000", groesse=9,
         mitte=False, zahl=False):
    cell.font = Font(name="Arial", bold=fett, color=farbe, size=groesse)
    cell.fill = PatternFill("solid", fgColor=bg)
    ha = "center" if mitte else ("right" if zahl else "left")
    cell.alignment = Alignment(horizontal=ha, vertical="center", wrap_text=False)
    cell.border = Border(
        bottom=Side(style="thin", color="DDDDDD"),
        right=Side(style="thin", color="DDDDDD"),
    )

def kopfzeile(ws, zeile, texte, bg=BLAU_D):
    for col, txt in enumerate(texte, 1):
        c = ws.cell(row=zeile, column=col, value=txt)
        stil(c, fett=True, bg=bg, farbe="FFFFFF", groesse=9, mitte=True)
    ws.row_dimensions[zeile].height = 22

def datenzeile(ws, zeile, werte, bg=WEISS):
    for col, wert in enumerate(werte, 1):
        c = ws.cell(row=zeile, column=col, value=wert)
        is_zahl = isinstance(wert, (int, float))
        stil(c, bg=bg, zahl=is_zahl)
        if is_zahl and col >= 4:
            c.number_format = '#,##0.00 €'

def breiten(ws, bmap):
    for col, b in bmap.items():
        ws.column_dimensions[col].width = b


# ── Blatt: Anfragen-Übersicht ─────────────────────────────────────────────────

def blatt_anfragen(wb, anfragen):
    ws = wb.create_sheet("Anfragen")
    ws.sheet_view.showGridLines = False

    # Titel
    ws["B1"].value = (
        f"CRM-Anfragen – Auftragswert < {LIMIT_KLEIN:,.0f} € oder > {LIMIT_GROSS:,.0f} €"
    )
    ws["B1"].font = Font(name="Arial", bold=True, size=13, color=BLAU_D)
    ws.row_dimensions[1].height = 30

    # Spalten dynamisch aus erstem Datensatz
    if not anfragen:
        ws["B3"].value = "(keine Anfragen gefunden)"
        return

    ignore = {"_positionen"}
    keys = [k for k in anfragen[0].keys() if k not in ignore]
    kopfzeile(ws, 3, keys)

    for i, a in enumerate(anfragen, 4):
        wert = a.get("Auftragswert") or 0
        try:
            wert_f = float(wert or 0)
        except (TypeError, ValueError):
            wert_f = 0
        bg = ROT_H if wert_f > LIMIT_GROSS else (BLAU_H if wert_f < LIMIT_KLEIN else WEISS)
        werte = [a.get(k) for k in keys]
        datenzeile(ws, i, werte, bg=bg)

    # Legende
    leg = len(anfragen) + 5
    ws.cell(row=leg, column=2, value="Rot = Großauftrag > 150 k €").font = Font(name="Arial", size=8, color="888888")
    ws.cell(row=leg+1, column=2, value="Blau = Kleinauftrag < 50 k €").font = Font(name="Arial", size=8, color="888888")

    # Auto-Filter + Einfrieren
    ws.auto_filter.ref = f"B3:{get_column_letter(len(keys)+1)}{len(anfragen)+3}"
    ws.freeze_panes = ws.cell(row=4, column=1)

    breiten(ws, {"A": 3, "B": 14, "C": 16, "D": 28, "E": 18, "F": 10, "G": 18})


# ── Blatt: Artikel-Positionen ─────────────────────────────────────────────────

def blatt_positionen(wb, anfragen, positionen):
    ws = wb.create_sheet("Artikel-Positionen")
    ws.sheet_view.showGridLines = False

    ws["B1"].value = "Angeforderte Artikel – alle Anfragen"
    ws["B1"].font = Font(name="Arial", bold=True, size=13, color=BLAU_D)
    ws.row_dimensions[1].height = 30

    if not positionen:
        ws["B3"].value = "(keine Positionen gefunden)"
        return

    # Alle Schlüssel aus erster Position
    keys = list(positionen[0].keys())
    kopfzeile(ws, 3, keys)

    for i, p in enumerate(positionen, 4):
        bg = BLAU_H if i % 2 == 0 else WEISS
        werte = [p.get(k) for k in keys]
        datenzeile(ws, i, werte, bg=bg)

    ws.auto_filter.ref = f"A3:{get_column_letter(len(keys))}{len(positionen)+3}"
    ws.freeze_panes = ws.cell(row=4, column=1)
    breiten(ws, {"A": 14, "B": 18, "C": 30, "D": 10, "E": 14, "F": 14, "G": 18, "H": 40})


# ── Blatt: Artikel-Zusammenfassung (dedupliziert) ─────────────────────────────

def blatt_artikel_summary(wb, positionen):
    ws = wb.create_sheet("Artikel-Zusammenfassung")
    ws.sheet_view.showGridLines = False

    ws["B1"].value = "Artikel-Zusammenfassung (eindeutig, nach Häufigkeit)"
    ws["B1"].font = Font(name="Arial", bold=True, size=13, color=BLAU_D)
    ws.row_dimensions[1].height = 30

    # Artikel aggregieren
    artikel_map = {}
    for p in positionen:
        artnr = str(p.get("Artikelnummer") or p.get("Bezeichnung") or "–")
        if artnr not in artikel_map:
            artikel_map[artnr] = {
                "Artikelnummer": p.get("Artikelnummer"),
                "Bezeichnung":   p.get("Bezeichnung"),
                "Anzahl_Anfragen": 0,
                "Gesamt_Menge":  0,
                "EK_Preis":      p.get("EK_Preis"),
                "VK_Preis":      p.get("VK_Preis"),
                "Langtext":      p.get("Langtext"),
            }
        artikel_map[artnr]["Anzahl_Anfragen"] += 1
        mng = p.get("Menge") or 0
        try:
            artikel_map[artnr]["Gesamt_Menge"] += float(mng)
        except (TypeError, ValueError):
            pass

    sortiert = sorted(artikel_map.values(),
                      key=lambda x: x["Anzahl_Anfragen"], reverse=True)

    headers = ["Artikelnummer", "Bezeichnung", "Anzahl Anfragen",
               "Gesamt-Menge", "EK-Preis €", "VK-Preis €",
               "✓ Korrekt?", "Anmerkung", "Langtext"]
    kopfzeile(ws, 3, headers)

    for i, art in enumerate(sortiert, 4):
        bg = GRUEN if i % 2 == 0 else WEISS
        werte = [
            art.get("Artikelnummer"),
            art.get("Bezeichnung"),
            art.get("Anzahl_Anfragen"),
            art.get("Gesamt_Menge"),
            art.get("EK_Preis"),
            art.get("VK_Preis"),
            "",   # Freigabe (manuell)
            "",   # Anmerkung (manuell)
            art.get("Langtext"),
        ]
        datenzeile(ws, i, werte, bg=bg)
        # Freigabe + Anmerkung gelb
        for col in [7, 8]:
            c = ws.cell(row=i, column=col)
            c.fill = PatternFill("solid", fgColor=GELB)

    ws.auto_filter.ref = f"A3:{get_column_letter(len(headers))}{len(sortiert)+3}"
    ws.freeze_panes = ws.cell(row=4, column=1)
    breiten(ws, {
        "A": 18, "B": 30, "C": 16, "D": 14,
        "E": 14, "F": 14, "G": 14, "H": 28, "I": 40,
    })


# ── Blatt: Freigabe ───────────────────────────────────────────────────────────

def blatt_freigabe(wb, heute, anzahl_anfragen, anzahl_artikel):
    ws = wb.create_sheet("Freigabe")
    ws.sheet_view.showGridLines = False

    ws["B2"].value = "Freigabe – CRM-Anfragen Stammdaten-Prüfung"
    ws["B2"].font = Font(name="Arial", bold=True, size=14, color=BLAU_D)
    ws.row_dimensions[2].height = 32

    ws["B3"].value = (
        f"Exportiert {heute.strftime('%d.%m.%Y')}  ·  "
        f"{anzahl_anfragen} Anfragen  ·  {anzahl_artikel} eindeutige Artikel"
    )
    ws["B3"].font = Font(name="Arial", size=9, color="888888")

    felder = [
        (6,  "Prüfer (Name)",                    ""),
        (7,  "Abteilung",                         "Vertrieb / Produktion"),
        (8,  "Datum der Prüfung",                 heute.strftime("%d.%m.%Y")),
        (9,  "Anfragen-Liste geprüft?",           "☐  Ja    ☐  Nein"),
        (10, "Artikel-Stammdaten vollständig?",   "☐  Ja    ☐  Nein"),
        (11, "Artikel-Zusammenfassung geprüft?",  "☐  Ja    ☐  Nein"),
        (13, "Gesamtfreigabe für Odoo-Import",    "☐  Freigegeben    ☐  Nicht freigegeben"),
        (15, "Unterschrift",                      ""),
    ]
    for zeile, label, vor in felder:
        ws.cell(row=zeile, column=2, value=label).font = Font(name="Arial", bold=True, size=10)
        c = ws.cell(row=zeile, column=3, value=vor)
        c.font = Font(name="Arial", size=10)
        c.fill = PatternFill("solid", fgColor=GELB)
        ws.row_dimensions[zeile].height = 24
        if "Unterschrift" in label:
            ws.row_dimensions[zeile].height = 40
            c.value = ""
            c.border = Border(bottom=Side(style="medium", color="333333"))

    ws["B17"].value = "Anmerkungen:"
    ws["B17"].font = Font(name="Arial", bold=True, size=10)
    ws.merge_cells("B18:F23")
    anm = ws["B18"]
    anm.fill = PatternFill("solid", fgColor=GELB)
    anm.alignment = Alignment(vertical="top", wrap_text=True)
    anm.border = Border(
        top=Side(style="thin"), bottom=Side(style="thin"),
        left=Side(style="thin"), right=Side(style="thin"),
    )
    for r in range(18, 24):
        ws.row_dimensions[r].height = 18

    breiten(ws, {"A": 3, "B": 38, "C": 50})


# ── Hauptprogramm ─────────────────────────────────────────────────────────────

def main():
    heute = date.today()
    conn  = ee_connect()

    info = erkunde_schema(conn)
    anfragen, positionen = lade_anfragen(conn, info)
    conn.close()

    if not anfragen:
        print("\n[!] Keine passenden Anfragen gefunden.")
        print("    Mögliche Ursachen:")
        print("    - Wertspalte nicht erkannt (wert_kopf = None)")
        print("    - Alle Anfragen liegen zwischen 50 k und 150 k €")
        print("    Bitte suche_pia_anfrage2.py Ausgabe für Spaltennamen prüfen.")
        return

    eindeutige_artikel = len({
        p.get("Artikelnummer") or p.get("Bezeichnung")
        for p in positionen
    })

    print(f"\n[3] Erstelle Excel …")
    wb = openpyxl.Workbook()
    wb.remove(wb.active)

    blatt_anfragen(wb, anfragen)
    blatt_positionen(wb, anfragen, positionen)
    blatt_artikel_summary(wb, positionen)
    blatt_freigabe(wb, heute, len(anfragen), eindeutige_artikel)

    datei = f"Anfragen_Pruefung_{heute.strftime('%Y%m%d')}.xlsx"
    wb.save(datei)
    print(f"\n[OK] Gespeichert: {datei}")
    print(f"     {len(anfragen)} Anfragen  ·  {len(positionen)} Positionen  ·  {eindeutige_artikel} eindeutige Artikel")
    print(f"     Bitte an Vertrieb / Produktionsleiter weiterleiten.")


if __name__ == "__main__":
    main()
